"""Analise funcional de todas as ferramentas do app.

Cobre cada handler do registry com arquivos reais e casos de borda,
alem de um teste end-to-end offscreen da OrganizePage.

Rode com:  python test_features.py
"""
import os
import sys
import tempfile
import time
import traceback
from pathlib import Path

os.environ.setdefault("QT_QPA_PLATFORM", "offscreen")

import pymupdf
from PySide6.QtWidgets import QApplication

from app.core import converters, pdf_tools
from app.registry import TOOLS

results = []


def check(name: str, ok: bool, detail: str = "") -> None:
    results.append((name, ok, detail))
    tag = "PASS" if ok else "FAIL"
    print(f"  [{tag}] {name}" + (f" — {detail}" if detail else ""))


def expect_error(name: str, fn, *args, **kwargs) -> None:
    try:
        fn(*args, **kwargs)
    except Exception as exc:
        check(name, True, f"erro esperado: {exc}")
        return
    check(name, False, "nao lancou erro")


def make_pdf(path: Path, labels: list[str]) -> None:
    """Cria PDF com uma pagina por label, texto unico por pagina."""
    doc = pymupdf.open()
    for label in labels:
        page = doc.new_page()
        page.insert_text((72, 300), label, fontsize=40)
    doc.save(path)
    doc.close()


def page_texts(path: Path) -> list[str]:
    with pymupdf.open(path) as doc:
        return [page.get_text().strip() for page in doc]


def main() -> None:
    app = QApplication(sys.argv)

    with tempfile.TemporaryDirectory() as tmp_s:
        tmp = Path(tmp_s)
        a = tmp / "A.pdf"
        b = tmp / "B.pdf"
        make_pdf(a, [f"A-pag{i}" for i in range(1, 6)])   # 5 paginas
        make_pdf(b, [f"B-pag{i}" for i in range(1, 4)])   # 3 paginas

        # imagens de exemplo
        img1, img2 = tmp / "img1.png", tmp / "img2.png"
        pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 100, 150))
        pix.clear_with(200)
        pix.save(img1)
        pix.save(img2)

        html = tmp / "pagina.html"
        html.write_text(
            "<html><head><style>"
            "h1{color:green} td{border:1px solid #333}"
            "</style></head><body>"
            "<h1>Titulo</h1><p>Conteudo de teste</p>"
            "<table><tr><td>A1</td><td>B1</td></tr></table>"
            "</body></html>",
            encoding="utf-8",
        )

        # PDF pesado com imagem ruidosa (dificil de comprimir) p/ compressao
        from PIL import Image
        noise_img = tmp / "ruido.png"
        Image.effect_noise((1500, 1500), 80).convert("RGB").save(noise_img)
        heavy = tmp / "pesado.pdf"
        hdoc = pymupdf.open()
        hpage = hdoc.new_page(width=595, height=842)
        hpage.insert_image(hpage.rect, filename=str(noise_img))
        hdoc.save(heavy)
        hdoc.close()

        # PDF com tabela com bordas p/ extracao de tabelas
        table_pdf = tmp / "tabela.pdf"
        tdoc = pymupdf.open()
        tpage = tdoc.new_page()
        data = [["Nome", "Idade", "Cidade"], ["Ana", "30", "SP"], ["Beto", "25", "RJ"]]
        for r, rowvals in enumerate(data):
            for c, val in enumerate(rowvals):
                rect = pymupdf.Rect(50 + c * 150, 50 + r * 30, 50 + (c + 1) * 150, 50 + (r + 1) * 30)
                tpage.draw_rect(rect)
                tpage.insert_text((rect.x0 + 5, rect.y0 + 20), val, fontsize=11)
        tdoc.save(table_pdf)
        tdoc.close()

        # ============================ ORGANIZAR ============================
        print("\n== ORGANIZAR ==")

        out = str(tmp / "merge.pdf")
        pdf_tools.merge_pdfs([str(a), str(b)], out)
        texts = page_texts(out)
        check(
            "Juntar: soma de paginas",
            len(texts) == 8,
            f"{len(texts)} paginas",
        )
        check(
            "Juntar: ordem preservada",
            texts[0].startswith("A-pag1") and texts[5].startswith("B-pag1"),
        )
        expect_error(
            "Juntar: erro com 1 arquivo",
            pdf_tools.merge_pdfs, [str(a)], str(tmp / "x.pdf"),
        )

        split_dir = tmp / "split"
        split_dir.mkdir()
        outs = pdf_tools.split_pdf([str(a)], str(split_dir))
        counts = [pymupdf.open(o).page_count for o in outs]
        check(
            "Dividir: 1 arquivo por pagina",
            len(outs) == 5 and all(c == 1 for c in counts),
            f"{len(outs)} arquivos",
        )
        text0 = page_texts(outs[0])[0]
        check("Dividir: conteudo correto", text0.startswith("A-pag1"))

        out = str(tmp / "remove.pdf")
        pdf_tools.remove_pages([str(a)], out, {"pages": "2-3"})
        texts = page_texts(out)
        check(
            "Remover: removeu paginas 2-3",
            len(texts) == 3 and texts[1].startswith("A-pag4"),
            f"restantes: {len(texts)}",
        )
        expect_error(
            "Remover: intervalo invalido 'abc'",
            pdf_tools.remove_pages, [str(a)], str(tmp / "x.pdf"), {"pages": "abc"},
        )
        expect_error(
            "Remover: fora do intervalo '9'",
            pdf_tools.remove_pages, [str(a)], str(tmp / "x.pdf"), {"pages": "9"},
        )
        expect_error(
            "Remover: campo vazio",
            pdf_tools.remove_pages, [str(a)], str(tmp / "x.pdf"), {"pages": ""},
        )

        out = str(tmp / "extract.pdf")
        pdf_tools.extract_pages([str(a)], out, {"pages": "1,4-5"})
        texts = page_texts(out)
        check(
            "Extrair: extraiu 1,4,5 na ordem",
            len(texts) == 3
            and texts[0].startswith("A-pag1")
            and texts[1].startswith("A-pag4")
            and texts[2].startswith("A-pag5"),
            f"{len(texts)} paginas",
        )
        # intervalo invertido "5-4" deve funcionar como "4-5"
        out2 = str(tmp / "extract2.pdf")
        pdf_tools.extract_pages([str(a)], out2, {"pages": "5-4"})
        check(
            "Extrair: intervalo invertido",
            len(page_texts(out2)) == 2,
        )
        expect_error(
            "Extrair: sem informar paginas",
            pdf_tools.extract_pages, [str(a)], str(tmp / "x.pdf"), None,
        )

        # ============================ OTIMIZAR =============================
        print("\n== OTIMIZAR ==")

        out = str(tmp / "compress.pdf")
        pdf_tools.compress_pdf([str(a)], out)
        check(
            "Comprimir: PDF valido, mesmas paginas",
            len(page_texts(out)) == 5,
            f"{Path(out).stat().st_size} bytes",
        )
        # PDF com imagem pesada deve realmente encolher
        out_heavy = str(tmp / "pesado_compress.pdf")
        pdf_tools.compress_pdf([str(heavy)], out_heavy)
        before, after = heavy.stat().st_size, Path(out_heavy).stat().st_size
        check(
            "Comprimir: reducao real de tamanho",
            after < before * 0.8,
            f"{before} -> {after} bytes ({100 * after // before}%)",
        )
        check(
            "Comprimir: PDF pesado continua abrindo",
            pymupdf.open(out_heavy).page_count == 1,
        )

        # corrompe o PDF (trunca o final) e tenta reparar
        broken = tmp / "broken.pdf"
        data = a.read_bytes()
        broken.write_bytes(data[:-200])
        out = str(tmp / "repaired.pdf")
        try:
            pdf_tools.repair_pdf([str(broken)], out)
            pages = pymupdf.open(out).page_count
            check("Reparar: PDF truncado recuperado", pages >= 1, f"{pages} paginas")
        except ValueError as exc:
            check("Reparar: PDF truncado", False, str(exc))

        # reparo de arquivo valido nao pode quebrar nada
        out = str(tmp / "repair_ok.pdf")
        pdf_tools.repair_pdf([str(a)], out)
        check("Reparar: arquivo sadio continua sadio", len(page_texts(out)) == 5)

        # ========================= CONVERTER EM PDF ========================
        print("\n== CONVERTER EM PDF ==")

        out = str(tmp / "imgs.pdf")
        converters.images_to_pdf([str(img1), str(img2)], out)
        check(
            "Imagem->PDF: 1 imagem por pagina",
            len(page_texts(out)) == 2,
        )

        out = str(tmp / "html.pdf")
        converters.html_to_pdf([str(html)], out)
        # xhtml2pdf (reportlab) gera texto selecionavel ate no modo offscreen
        texts = page_texts(out)
        check(
            "HTML->PDF: texto extraivel no PDF",
            len(texts) >= 1 and "Titulo" in "".join(texts),
            f"{len(texts)} pagina(s)",
        )

        # ========================= CONVERTER DE PDF ========================
        print("\n== CONVERTER DE PDF ==")

        jpg_dir = tmp / "jpgs"
        jpg_dir.mkdir()
        outs = converters.pdf_to_jpg([str(a)], str(jpg_dir))
        valid = all(
            Path(o).read_bytes()[:2] == b"\xff\xd8"  # magic JPEG
            for o in outs
        )
        check(
            "PDF->JPG: 1 jpg valido por pagina",
            len(outs) == 5 and valid,
            f"{len(outs)} imagens",
        )

        out = str(tmp / "doc.docx")
        converters.pdf_to_word([str(a)], out)
        check(
            "PDF->Word: .docx gerado",
            Path(out).stat().st_size > 1000
            and Path(out).read_bytes()[:2] == b"PK",  # zip = docx
            f"{Path(out).stat().st_size} bytes",
        )

        out = str(tmp / "slides.pptx")
        converters.pdf_to_pptx([str(a)], out)
        from pptx import Presentation
        slides = len(Presentation(out).slides.__iter__.__self__._sldIdLst)
        check(
            "PDF->PowerPoint: 1 slide por pagina",
            slides == 5,
            f"{slides} slides",
        )

        out = str(tmp / "planilha.xlsx")
        converters.pdf_to_excel([str(table_pdf)], out)
        from openpyxl import load_workbook
        wb = load_workbook(out)
        sheet = wb.active
        cells = {
            (c.row, c.column): c.value
            for row in sheet.iter_rows()
            for c in row
            if c.value is not None
        }
        check(
            "PDF->Excel: tabela em colunas",
            cells.get((1, 1)) == "Nome"
            and cells.get((1, 2)) == "Idade"
            and cells.get((2, 2)) == "30",
            str(cells),
        )
        # fallback: PDF sem tabela vira linhas de texto
        out2 = str(tmp / "planilha2.xlsx")
        converters.pdf_to_excel([str(a)], out2)
        wb2 = load_workbook(out2)
        check(
            "PDF->Excel: fallback texto (1 aba/pagina)",
            len(wb2.sheetnames) == 5,
            f"abas: {wb2.sheetnames}",
        )

        # ==================== ORGANIZAR (end-to-end UI) ====================
        print("\n== ORGANIZAR (UI offscreen) ==")
        from unittest.mock import patch
        from app.ui.pages.organize_page import OrganizePage
        from app.registry import get_tool

        page = OrganizePage(get_tool("organize"))
        page._load([str(a), str(b)])
        check(
            "Organizar: carrega multiplos PDFs",
            page.pages.count() == 8 and len(page._docs) == 2,
        )

        # move item 0 (A-pag1) para depois do item 1 -> [A2, A1, A3, ...]
        page.pages._reinsert([0], 2)
        order = page._order()
        check(
            "Organizar: reordenacao",
            order[:3] == [(0, 1), (0, 0), (0, 2)],
            str(order[:3]),
        )

        # gira a primeira pagina
        page.pages.item(0).setSelected(True)
        page._rotate_selected()
        check(
            "Organizar: rotacao registrada",
            page._rotations.get((0, 1)) == 90,
        )

        # remove a segunda
        page.pages.clearSelection()
        page.pages.item(1).setSelected(True)
        page._remove_selected()
        check(
            "Organizar: remocao",
            page.pages.count() == 7,
        )

        # salva e confere o arquivo final
        out = str(tmp / "organizado.pdf")
        with patch(
            "app.ui.pages.organize_page.QFileDialog.getSaveFileName",
            return_value=(out, ""),
        ):
            page._save()
        with pymupdf.open(out) as doc:
            saved_pages, saved_rot = doc.page_count, doc[0].rotation
        check(
            "Organizar: salva ordem+rotacao+remocao",
            saved_pages == 7 and saved_rot == 90,
            f"{saved_pages} paginas, rotacao {saved_rot}",
        )

        # ======================= Organizar em lote =========================
        print("\n== ORGANIZAR EM LOTE ==")
        from app.core import auto_sort

        root = tmp / "SEMANA"
        # PDFs com texto digital (>=200 chars) pulam o OCR e classificam
        # direto pelas palavras-chave do perfil.
        filler = "palavra de preenchimento " * 12
        pac1 = root / "PACIENTE UM"
        pac1.mkdir(parents=True)
        doc = pymupdf.open()
        pg = doc.new_page()
        pg.insert_text((72, 300), f"ESPELHO DA AIH\n{filler}", fontsize=12)
        pg = doc.new_page()
        pg.insert_text(
            (72, 300), f"GUIA DE AUTORIZACAO DE INTERNAMENTO\n{filler}",
            fontsize=12,
        )
        doc.save(pac1 / "scan1.pdf")
        doc.close()
        (root / "PASTA VAZIA").mkdir()

        batch_results = auto_sort.organize_batch(root)
        check(
            "Lote: 2 pastas processadas",
            len(batch_results) == 2,
            str([r["status"] for r in batch_results]),
        )
        # faltam AIH, FATURAMENTO e DANFE -> a anotacao vai no nome
        # do dossie E da pasta (renomeada)
        pac1_ok = next(root.glob("PACIENTE UM*"))
        saida = next(
            p for p in pac1_ok.glob("*.pdf")
            if not p.name.endswith(".tmp.pdf")
            and p.stem.startswith("PACIENTE UM")
        )
        check(
            "Lote: saida anota docs faltantes no nome",
            "falta" in saida.stem and "(falta" in pac1_ok.name
            and "AIH" in saida.stem,
            saida.name,
        )
        check("Lote: saida salva na pasta do paciente", saida.exists())
        if saida.exists():
            texts = page_texts(saida)
            check(
                "Lote: guia vem antes do espelho",
                len(texts) == 2 and "GUIA" in texts[0] and "ESPELHO" in texts[1],
            )
        # segunda rodada: pasta ja organizada deve ser pulada
        rerun = auto_sort.organize_batch(root)
        check(
            "Lote: rerun pula pasta ja organizada",
            rerun[0]["status"] == "skip",
            rerun[0].get("detail", ""),
        )
        # crash simulation: saida truncada + resto .tmp.pdf -> reprocessa
        saida.write_bytes(b"%PDF-1.4 truncado")  # corrompe a saida
        tmp_crash = saida.with_suffix(".tmp.pdf")
        tmp_crash.write_bytes(b"lixo de crash")
        rerun = auto_sort.organize_batch(root)
        ok_rerun = rerun[0]
        check(
            "Lote: saida corrompida e reprocessada",
            ok_rerun["status"] == "ok"
            and auto_sort._is_valid_output(saida)
            and not tmp_crash.exists(),
            f"{ok_rerun['status']}",
        )
        # auditoria: guia com nome de outro paciente deve ser suspeita
        pac2 = root / "MARIA SILVA"
        pac2.mkdir()
        doc = pymupdf.open()
        pg = doc.new_page()
        pg.insert_text(
            (72, 300),
            f"GUIA DE AUTORIZACAO DE INTERNAMENTO\n"
            f"Sr(a), JOSE DOS SANTOS FILHO - Idade: 40a.\n{filler}",
            fontsize=12,
        )
        doc.save(pac2 / "scan.pdf")
        doc.close()
        with pymupdf.open(pac2 / "scan.pdf") as d2:
            audit = auto_sort.audit_patient(
                [d2], auto_sort.ref_from_folder("MARIA SILVA")
            )
        check(
            "Auditoria: pagina de outro paciente e suspeita",
            len(audit["suspects"]) == 1
            and audit["suspects"][0].get("found"),
            str(audit["suspects"]),
        )

        # pagina anulada (X sobre o documento) e excluida do dossie
        pac3 = root / "JOSE LIMA"
        pac3.mkdir()
        doc = pymupdf.open()
        doc.new_page()                      # p1: pagina normal
        pg = doc.new_page()                 # p2: anulada com X
        r = pg.rect
        pg.draw_line(r.tl, r.br, color=(0, 0, 0), width=5)
        pg.draw_line(r.tr, r.bl, color=(0, 0, 0), width=5)
        doc.save(pac3 / "scan.pdf")
        doc.close()
        with pymupdf.open(pac3 / "scan.pdf") as d3:
            cls = auto_sort.classify_documents([d3])
            order = auto_sort.suggest_order([d3])
        check(
            "Anulado: pagina com X e excluida",
            cls.get((0, 1)) == "anulado"
            and cls.get((0, 0)) != "anulado"
            and order == [(0, 0)],
            f"{cls} / {order}",
        )

        # arquivo corrompido no meio nao pode derrubar a pasta —
        # os PDFs legiveis ainda entram no dossie e o ilegivel e
        # reportado em "skipped"
        pac4 = root / "ANA SOUZA"
        pac4.mkdir()
        doc = pymupdf.open()
        pg = doc.new_page()
        pg.insert_text(
            (72, 300), f"ESPELHO DA AIH\n{filler}", fontsize=12
        )
        doc.save(pac4 / "bom.pdf")
        doc.close()
        (pac4 / "corrompido.pdf").write_bytes(b"%PDF-1.4 lixo")
        rerun4 = auto_sort.organize_batch(root)
        entry4 = next(r for r in rerun4 if r["folder"] == "ANA SOUZA")
        check(
            "Lote: arquivo ilegivel pulado, pasta processada",
            entry4["status"] == "ok"
            and entry4.get("skipped") == ["corrompido.pdf"]
            and entry4["pages"] == 1,
            str(entry4),
        )
        # pasta so com arquivo corrompido vira erro claro
        pac5 = root / "SEM PDF VALIDO"
        pac5.mkdir()
        (pac5 / "ruim.pdf").write_bytes(b"nao e pdf")
        rerun5 = auto_sort.organize_batch(root)
        entry5 = next(r for r in rerun5 if r["folder"] == "SEM PDF VALIDO")
        check(
            "Lote: pasta so com ilegiveis vira erro",
            entry5["status"] == "error"
            and "legivel" in entry5.get("detail", ""),
            str(entry5),
        )

        # cache de OCR: arquivo substituido (mesmo nome) nao reusa
        # resultado velho — a assinatura inclui mtime/tamanho
        pac4_ok = next(root.glob("ANA SOUZA*"))  # pasta foi renomeada
        with pymupdf.open(pac4_ok / "bom.pdf") as d_sig:
            sig1 = auto_sort._file_sig(d_sig, "x")
        time.sleep(0.02)
        doc = pymupdf.open()
        doc.new_page()
        doc.save(pac4_ok / "bom.pdf")  # sobrescreve: outro conteudo
        doc.close()
        with pymupdf.open(pac4_ok / "bom.pdf") as d_sig:
            sig2 = auto_sort._file_sig(d_sig, "x")
        check(
            "Cache: assinatura muda quando arquivo e substituido",
            sig1 != sig2,
            f"{sig1} vs {sig2}",
        )

        # pasta onde TODAS as paginas sao anuladas vira erro — um
        # PDF vazio salvo seria invalido e reprocessaria toda rodada
        pac6 = root / "SO ANULADO"
        pac6.mkdir()
        doc = pymupdf.open()
        pg = doc.new_page()
        r = pg.rect
        pg.draw_line(r.tl, r.br, color=(0, 0, 0), width=5)
        pg.draw_line(r.tr, r.bl, color=(0, 0, 0), width=5)
        doc.save(pac6 / "scan.pdf")
        doc.close()
        rerun6 = auto_sort.organize_batch(root)
        entry6 = next(r for r in rerun6 if r["folder"] == "SO ANULADO")
        check(
            "Lote: pasta so com anuladas vira erro",
            entry6["status"] == "error"
            and "anulada" in entry6.get("detail", ""),
            str(entry6),
        )
        # PDF solto na raiz (ao lado das subpastas) nao e ignorado
        doc = pymupdf.open()
        doc.new_page()
        doc.save(root / "solto.pdf")
        doc.close()
        folders_now = auto_sort.list_patient_folders(root)
        check(
            "Lote: PDF solto na raiz vira pseudo-pasta",
            folders_now[0][0] == root
            and any(p.name == "solto.pdf" for p in folders_now[0][1]),
            str([f.name for f, _ in folders_now]),
        )
        # a saida <pasta>.pdf (mesmo com anotacao falta) nao conta
        # como insumo da pasta
        pac1_key = next(
            f for f, _ in folders_now if f.name.startswith("PACIENTE UM")
        )
        pac1_pdfs = dict(folders_now)[pac1_key]
        check(
            "Lote: saida nao conta como arquivo de entrada",
            all(
                "falta" not in p.stem and p.stem != "PACIENTE UM"
                for p in pac1_pdfs
            ),
            str([p.name for p in pac1_pdfs]),
        )

        # pagina da UI monta e lista pastas
        from app.ui.pages.batch_page import BatchPage
        from app.registry import get_tool
        bpage = BatchPage(get_tool("lote_ortopedia"))
        bpage._scan(root)
        check(
            "Lote: pagina lista as pastas",
            bpage.folder_list.count() == 8 and bpage.run_btn.isEnabled(),
        )

        # Organizar: mesmo PDF duas vezes nao duplica a grade
        from app.ui.pages.organize_page import OrganizePage
        opage = OrganizePage(get_tool("organize"))
        opage._load([str(a), str(a), str(b)])
        check(
            "Organizar: dedup de arquivos repetidos",
            opage.pages.count() == 8 and len(opage._docs) == 2,
            f"{opage.pages.count()} paginas, {len(opage._docs)} docs",
        )
        opage._close_docs()

        # ======================= Janela principal ==========================
        print("\n== JANELA PRINCIPAL ==")
        from app.ui.main_window import MainWindow
        win = MainWindow()
        check(
            "MainWindow: todas as paginas montam",
            True,
            f"{len(TOOLS)} ferramentas registradas",
        )
        # toda tool com handler tem handler chamavel; organize e excecao
        for t in TOOLS:
            if t.id not in ("organize", "lote_ortopedia") and not callable(t.handler):
                check(f"handler {t.id}", False, "handler None")

        page._close_docs()  # libera os PDFs abertos para o temp dir fechar

    # ------------------------------------------------------------ resumo
    fails = [r for r in results if not r[1]]
    print(f"\n{'='*60}")
    print(f"RESULTADO: {len(results) - len(fails)}/{len(results)} testes passaram")
    if fails:
        print("FALHAS:")
        for name, _, detail in fails:
            print(f"  - {name}: {detail}")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
